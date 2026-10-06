#!/usr/bin/env python3
"""질문별 Agent 답변 현황 → Excel 리포트 (운영계 실행용)

입력  = 90_provided_definition/데이터플랫폼 ML예측프롬프트_sample.csv (전체 질문)
호출  = SNOWFLAKE.CORTEX.DATA_AGENT_RUN('GN_DW.SERVING.<AGENT>', 요청 JSON)
출력  = 12_agent개선과제/output/질문별_답변현황_<계정>_<YYYYMMDD_HHMM>.xlsx
        + output/raw/<태그>.json (응답 원문 · 재실행 시 캐시로 재사용)

실행 환경 (둘 다 지원)
  ㉠ CoCo 샌드박스(bash)  : python3 "12_agent개선과제/tools/agent_answer_report.py" --apply
     접속 = SNOWFLAKE_ACCOUNT · SNOWFLAKE_HOST · SNOWFLAKE_TOKEN_FILE_PATH 환경변수(OAuth 토큰)
  ㉡ Workspace Python 파일 : Snowpark get_active_session() 이 있으면 그것을 쓴다(병렬 1로 고정)

🔴 LLM 대량 호출(크레딧 과금)이다 — `--apply` 가 없으면 질문 목록만 출력하고 끝낸다.
🔴 답변 원문은 화면에 출력하지 않는다(파일로만 남긴다).

주요 옵션
  --apply              실제 호출
  --ym YYYYMM          자리표시 기준년월(기본 = 실행일 전월)
  --only 1,3,10B       CSV 행번호(1부터 · B안은 10B 처럼) 일부만
  --agents A,B         부서 배정 대신 지정 Agent 로만 호출
  --workers N          동시 호출 수(기본 4 · Snowpark 세션이면 1)
  --role R --warehouse W   접속 역할·웨어하우스(기본 = 환경변수 또는 세션 현재값)
  --force              캐시(raw/*.json)를 무시하고 다시 호출
  --report-only        호출 없이 raw 캐시로 Excel 만 다시 만든다(과금 0)
  --out DIR            출력 폴더(기본 = 12_agent개선과제/output)
"""
import argparse
import csv
import datetime as dt
import io
import json
import os
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
CSV_PATH = os.path.join(ROOT, '90_provided_definition', '데이터플랫폼 ML예측프롬프트_sample.csv')
OUT_DIR = os.path.join(os.path.dirname(HERE), 'output')
RAW_DIR = os.path.join(OUT_DIR, 'raw')
AGENT_DB_SCHEMA = 'GN_DW.SERVING'

# 부서 → Agent (O202-C `부서별agent답변.md` 와 같은 배정)
DEPT_AGENT = {
    '회원실': ['AGENT_MEMBER'],
    '기획실': ['AGENT_MEMBER'],
    '나눔마케팅본부': ['AGENT_MARKETING'],
}
DEFAULT_AGENT = ['AGENT_MEMBER']
# 같은 질문을 추가 Agent 에도 묻는 행(CSV 행번호) — LTV 채널·캠페인(종전 #29·#31)
EXTRA_AGENTS = {32: ['AGENT_EXECUTIVE'], 34: ['AGENT_EXECUTIVE']}

EXCEL_CELL_MAX = 32000
ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


# ---------------------------------------------------------------- 질문 구성
def prev_month(today):
    first = today.replace(day=1)
    return (first - dt.timedelta(days=1)).strftime('%Y%m')


def fill(text, ym):
    y, m = int(ym[:4]), int(ym[4:])
    last = (dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).isoformat()
    rep = {
        '{기준년월}': f'{y}-{m:02d}', '{YYYY-MM}': f'{y}-{m:02d}', '{기준월}': f'{y}-{m:02d}',
        '{YYYY-MM-DD}': last, '{기준연도}': str(y), '{연도}': str(y), '{차년도}': str(y + 1),
    }
    for k, v in rep.items():
        text = text.replace(k, v)
    # 그 밖의 자리표시({10월~12월} · {당해년도 12월말 활동회원 예상치} 등)는 괄호만 벗긴다
    return re.sub(r'\{([^}]+)\}', r'\1', text)


def load_questions(ym):
    rows = list(csv.reader(io.open(CSV_PATH, encoding='utf-8-sig')))[1:]
    out, dept, kind, sub, prev = [], '', '', '', None
    for i, r in enumerate(rows, 1):
        r = (r + [''] * 7)[:7]
        dept = r[0].strip() or dept
        kind = r[1].strip() or kind
        if r[3].strip():
            sub = r[2].strip() or sub
        prompt, need, detail, note = r[3].strip(), r[4].strip(), r[5].strip(), r[6].strip()
        variant, qrow = '', i
        if not prompt and prev is not None and detail:      # B안 행 = 직전 질문의 변형(같은 번호)
            prompt, need, variant, qrow = prev['prompt'], prev['need'], 'B', prev['row']
            if not prev['variant']:
                prev['variant'] = 'A'
                prev['key'] = f"{prev['row']}A"
        elif not prompt:
            continue
        key = f'{qrow}{variant}'
        parts = [fill(prompt, ym)]
        if need:
            parts.append('[필수 조건값]\n' + fill(need, ym))
        if detail:
            parts.append('[세부 조건값]\n' + fill(detail, ym))
        if note and dept == '기획실':                       # 기획실 = 산출 방식(A안·B안 수식)
            parts.append('[산출 방식]\n' + fill(note, ym))
        q = dict(row=qrow, csv_row=i, key=key, dept=dept, kind=kind, sub=sub, variant=variant,
                 prompt=prompt, need=need, detail=detail, note=note,
                 question='\n\n'.join(parts))
        out.append(q)
        if r[3].strip():
            prev = q
    return out


def plan(questions, agents_override):
    jobs = []
    for q in questions:
        agents = agents_override or (DEPT_AGENT.get(q['dept'], DEFAULT_AGENT)
                                     + EXTRA_AGENTS.get(q['row'], []))
        for a in agents:
            jobs.append(dict(q, agent=a, tag=f"Q{q['row']:02d}{q['variant']}_{a}"))
    return jobs


# ---------------------------------------------------------------- 접속
class Runner:
    def __init__(self, role, warehouse):
        self.role, self.warehouse = role, warehouse
        self.session = None
        try:
            from snowflake.snowpark.context import get_active_session
            self.session = get_active_session()
        except Exception:
            self.session = None
        self._local = threading.local()

    @property
    def snowpark(self):
        return self.session is not None

    def _conn(self):
        cn = getattr(self._local, 'cn', None)
        if cn is None:
            import snowflake.connector
            tok = open(os.environ.get('SNOWFLAKE_TOKEN_FILE_PATH',
                                      '/snowflake/session/token')).read().strip()
            kw = dict(account=os.environ['SNOWFLAKE_ACCOUNT'], host=os.environ.get('SNOWFLAKE_HOST'),
                      token=tok, authenticator='oauth', client_session_keep_alive=True)
            role = self.role or os.environ.get('SNOWFLAKE_ROLE')
            wh = self.warehouse or os.environ.get('SNOWFLAKE_WAREHOUSE')
            if role:
                kw['role'] = role
            if wh:
                kw['warehouse'] = wh
            cn = snowflake.connector.connect(**kw)
            self._local.cn = cn
        return cn

    def query(self, sql, params=None):
        if self.snowpark:
            return [tuple(r) for r in self.session.sql(sql, params=params).collect()]
        cur = self._conn().cursor()
        cur.execute(sql, params)
        return cur.fetchall()

    def run_agent(self, agent, question):
        req = json.dumps({'messages': [{'role': 'user',
                                        'content': [{'type': 'text', 'text': question}]}]},
                         ensure_ascii=False)
        return self.query('SELECT SNOWFLAKE.CORTEX.DATA_AGENT_RUN(?, ?)'
                          if self.snowpark else
                          'SELECT SNOWFLAKE.CORTEX.DATA_AGENT_RUN(%s, %s)',
                          [f'{AGENT_DB_SCHEMA}.{agent}', req])[0][0]


# ---------------------------------------------------------------- 응답 해석
def table_text(rs, max_rows=30):
    try:
        cols = [c.get('name') for c in rs['resultSetMetaData']['rowType']]
        data = rs.get('data') or []
    except Exception:
        return ''
    lines = [' | '.join(map(str, cols))]
    for r in data[:max_rows]:
        lines.append(' | '.join('' if v is None else str(v) for v in r))
    if len(data) > max_rows:
        lines.append(f'… 외 {len(data) - max_rows}행')
    return '\n'.join(lines)


def parse(raw):
    d = dict(answer='', tools=[], svs=[], sqls=[], tables=0, errors=[], status='',
             tokens=None, query_ids=[])
    try:
        j = json.loads(raw) if isinstance(raw, str) else raw
    except Exception:
        d['answer'] = str(raw)
        d['errors'].append('응답 JSON 해석 실패')
        return d
    d['status'] = j.get('status', '')
    try:
        tk = j['metadata']['usage']['tokens_consumed']
        d['tokens'] = sum(t['input_tokens']['total'] + t['output_tokens']['total'] for t in tk)
    except Exception:
        pass
    texts = []
    for it in j.get('content', []):
        t = it.get('type')
        v = it.get(t) if isinstance(it.get(t), dict) else {}
        if t == 'text':
            texts.append(v.get('text', it.get('text', '')) if v else it.get('text', ''))
        elif t == 'tool_use':
            name = v.get('name')
            if name and name not in d['tools']:
                d['tools'].append(name)
        elif t == 'tool_result':
            for c in v.get('content', []):
                js = c.get('json') or {}
                if js.get('semantic_view_fqn') and js['semantic_view_fqn'] not in d['svs']:
                    d['svs'].append(js['semantic_view_fqn'])
                if js.get('sql'):
                    d['sqls'].append(js['sql'])
                if js.get('query_id'):
                    d['query_ids'].append(js['query_id'])
                for k in ('error', 'message'):
                    if isinstance(js.get(k), str) and re.search(r'error|fail|invalid', js[k], re.I):
                        d['errors'].append(js[k][:300])
        elif t == 'table':
            d['tables'] += 1
            title = v.get('title') or f"표 {d['tables']}"
            texts.append(f'\n[{title}]\n' + table_text(v.get('result_set') or {}) + '\n')
    d['answer'] = ''.join(texts).strip()
    return d


NO_ANSWER = re.compile(r'답변할 수 없|제공하지 않|조회되지 않|예측(?:은|이) 없|모델(?:은|이) 없|불가합니다|불가능합니다')
PARTIAL = re.compile(r'대신|대체|대안|참고치|추세 참고|직접 예측(?:은|이) 아니|대리지표|근사')


def auto_judge(d, ok):
    """자동 판정 후보 — 사람이 「판정(확정)」 열에서 확정한다."""
    if not ok:
        return '호출실패'
    if not d['sqls']:
        return '✕ 후보'
    if NO_ANSWER.search(d['answer']) or PARTIAL.search(d['answer']):
        return '△ 후보'
    return '⭕ 후보'


# ---------------------------------------------------------------- 실행
def execute(runner, jobs, force, workers):
    os.makedirs(RAW_DIR, exist_ok=True)
    results = {}

    def one(job):
        path = os.path.join(RAW_DIR, job['tag'] + '.json')
        if os.path.exists(path) and not force:
            rec = json.load(io.open(path, encoding='utf-8'))
            rec['cached'] = True
            return job['tag'], rec
        t0 = time.time()
        try:
            raw, ok, err = runner.run_agent(job['agent'], job['question']), True, ''
        except Exception as e:
            raw, ok, err = '', False, f'{type(e).__name__}: {e}'
        rec = dict(tag=job['tag'], agent=job['agent'], ok=ok, error=err,
                   sec=round(time.time() - t0, 1), raw=raw,
                   called_at=dt.datetime.now().isoformat(timespec='seconds'))
        with io.open(path, 'w', encoding='utf-8') as f:
            json.dump(rec, f, ensure_ascii=False)
        return job['tag'], rec

    n = len(jobs)
    if workers <= 1:
        for k, job in enumerate(jobs, 1):
            tag, rec = one(job)
            results[tag] = rec
            print(f'[{k}/{n}] {tag} ok={rec["ok"]} {rec["sec"]}s'
                  + (' (캐시)' if rec.get('cached') else ''), flush=True)
    else:
        with ThreadPoolExecutor(max_workers=workers) as ex:
            futs = {ex.submit(one, j): j for j in jobs}
            for k, f in enumerate(as_completed(futs), 1):
                tag, rec = f.result()
                results[tag] = rec
                print(f'[{k}/{n}] {tag} ok={rec["ok"]} {rec["sec"]}s'
                      + (' (캐시)' if rec.get('cached') else ''), flush=True)
    return results


def load_cache(jobs):
    res = {}
    for j in jobs:
        p = os.path.join(RAW_DIR, j['tag'] + '.json')
        if os.path.exists(p):
            res[j['tag']] = json.load(io.open(p, encoding='utf-8'))
    return res


def clean(v):
    s = '' if v is None else str(v)
    s = ILLEGAL.sub('', s)
    return s if len(s) <= EXCEL_CELL_MAX else s[:EXCEL_CELL_MAX] + '\n…(이하 생략 · raw 파일 참조)'


def write_excel(jobs, results, info):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = '답변현황'
    head = ['No', '태그', '요청부서', '구분', '세부구분', '안', '원 프롬프트', '전송 질문',
            'Agent', '호출성공', '소요(초)', '사용 도구', '사용 SV', 'SQL 수', '실행 SQL',
            '표 수', '답변', '자동판정(후보)', '판정(확정)', '검토 의견', '오류', '토큰', 'query_id', '호출시각']
    ws.append(head)
    rows = []
    for k, j in enumerate(jobs, 1):
        rec = results.get(j['tag'])
        if rec is None:
            d, ok, sec, err, at = parse('{}'), None, None, '미호출', ''
        else:
            ok, sec, err, at = rec['ok'], rec['sec'], rec.get('error', ''), rec.get('called_at', '')
            d = parse(rec['raw']) if ok else parse('{}')
        judge = auto_judge(d, ok) if ok is not None else '미호출'
        err_all = '\n'.join([e for e in [err] + d['errors'] if e])
        row = [k, j['tag'], j['dept'], j['kind'], j['sub'], j['variant'], j['prompt'], j['question'],
               j['agent'], {True: 'Y', False: 'N', None: ''}[ok], sec,
               ', '.join(d['tools']), ', '.join(d['svs']), len(d['sqls']),
               '\n\n'.join(d['sqls'][:3]), d['tables'], d['answer'], judge, '', '',
               err_all, d['tokens'], ', '.join(d['query_ids'][:5]), at]
        ws.append([clean(v) if isinstance(v, str) else v for v in row])
        rows.append(dict(dept=j['dept'], agent=j['agent'], ok=ok, judge=judge, sec=sec or 0))

    widths = {'A': 5, 'B': 24, 'C': 12, 'D': 12, 'E': 16, 'F': 4, 'G': 40, 'H': 50, 'I': 18,
              'J': 8, 'K': 8, 'L': 30, 'M': 34, 'N': 7, 'O': 50, 'P': 6, 'Q': 90, 'R': 12,
              'S': 10, 'T': 30, 'U': 30, 'V': 10, 'W': 30, 'X': 18}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    hdr_fill = PatternFill('solid', fgColor='1F4E78')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(head))}{ws.max_row}'
    dv = DataValidation(type='list', formula1='"⭕,△,✕"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'S2:S{ws.max_row}')

    # 부서 요약
    s = wb.create_sheet('부서요약')
    cats = ['⭕ 후보', '△ 후보', '✕ 후보', '호출실패', '미호출']
    s.append(['요청부서', 'Agent', '호출', '성공'] + cats + ['평균 소요(초)'])
    keys = []
    for r in rows:
        if (r['dept'], r['agent']) not in keys:
            keys.append((r['dept'], r['agent']))
    for dept, agent in keys:
        g = [r for r in rows if r['dept'] == dept and r['agent'] == agent]
        okn = sum(1 for r in g if r['ok'])
        avg = round(sum(r['sec'] for r in g) / len(g), 1) if g else 0
        s.append([dept, agent, len(g), okn] + [sum(1 for r in g if r['judge'] == c) for c in cats] + [avg])
    s.append(['합계', '', len(rows), sum(1 for r in rows if r['ok'])]
             + [sum(1 for r in rows if r['judge'] == c) for c in cats] + [''])
    s.append([])
    s.append(['※ 자동판정은 후보다 — SQL 실행 없음 = ✕ 후보 · 대체/불가 문구 = △ 후보 · 그 외 = ⭕ 후보.'])
    s.append(["※ 확정 판정은 '답변현황' 시트의 「판정(확정)」 열(⭕/△/✕)에 사람이 입력한다."])
    for c in s[1]:
        c.font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJ', [16, 18, 6, 6, 9, 9, 9, 9, 9, 12]):
        s.column_dimensions[col].width = w

    # 실행정보
    m = wb.create_sheet('실행정보')
    for k, v in info.items():
        m.append([k, clean(v)])
    m.column_dimensions['A'].width = 22
    m.column_dimensions['B'].width = 100

    os.makedirs(OUT_DIR, exist_ok=True)
    stamp = dt.datetime.now().strftime('%Y%m%d_%H%M')
    path = os.path.join(OUT_DIR, f"질문별_답변현황_{info.get('계정', 'NA')}_{stamp}.xlsx")
    tmp_path = f"/tmp/report_{stamp}.xlsx"
    wb.save(tmp_path)
    with open(tmp_path, 'rb') as f_in, open(path, 'wb') as f_out:
        f_out.write(f_in.read())
    return path


def agent_versions(runner, agents):
    out = {}
    for a in agents:
        try:
            runner.query(f'SHOW VERSIONS IN AGENT {AGENT_DB_SCHEMA}.{a}')
            rs = runner.query('SELECT "name" FROM TABLE(RESULT_SCAN(LAST_QUERY_ID())) '
                              'WHERE "is_default" = \'true\'')
            out[a] = rs[0][0] if rs else '?'
        except Exception as e:
            out[a] = f'조회 실패: {type(e).__name__}'
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--apply', action='store_true')
    ap.add_argument('--ym')
    ap.add_argument('--only')
    ap.add_argument('--agents')
    ap.add_argument('--workers', type=int, default=4)
    ap.add_argument('--role')
    ap.add_argument('--warehouse')
    ap.add_argument('--force', action='store_true')
    ap.add_argument('--report-only', action='store_true')
    ap.add_argument('--out', help='출력 폴더(기본 = 12_agent개선과제/output)')
    argv = sys.argv[1:]
    if not argv and os.environ.get('AGENT_REPORT_ARGS'):   # 인자를 못 넘기는 실행기용
        import shlex
        argv = shlex.split(os.environ['AGENT_REPORT_ARGS'])
    a = ap.parse_args(argv)
    if a.out:
        global OUT_DIR, RAW_DIR
        OUT_DIR = os.path.abspath(a.out)
        RAW_DIR = os.path.join(OUT_DIR, 'raw')

    ym = a.ym or prev_month(dt.date.today())
    if not re.fullmatch(r'\d{6}', ym):
        sys.exit('--ym 은 YYYYMM 형식이다')
    qs = load_questions(ym)
    if a.only:
        want = {x.strip().upper() for x in a.only.split(',')}
        qs = [q for q in qs if q['key'].upper() in want or str(q['row']) in want]
    jobs = plan(qs, [x.strip() for x in a.agents.split(',')] if a.agents else None)

    print(f'질문 {len(qs)}개 · 호출 {len(jobs)}회 · 기준년월 {ym}')
    if not a.apply and not a.report_only:
        for j in jobs:
            print(f"  {j['tag']:<28} {j['dept']:<8} {j['question'][:60]!r}")
        print('\n드라이런이다(과금 0). 실제 호출 = --apply · 캐시로 Excel 만 = --report-only')
        return 0

    if a.report_only:
        runner, results = None, load_cache(jobs)
        info_extra = {'계정': 'CACHE', '실행 방식': 'report-only(캐시)'}
    else:
        runner = Runner(a.role, a.warehouse)
        workers = 1 if runner.snowpark else max(1, a.workers)
        acct, role, wh = runner.query('SELECT CURRENT_ACCOUNT(), CURRENT_ROLE(), CURRENT_WAREHOUSE()')[0]
        print(f'접속 = {acct} · {role} · {wh} · ' + ('Snowpark 세션' if runner.snowpark else 'connector'))
        results = execute(runner, jobs, a.force, workers)
        vers = agent_versions(runner, sorted({j['agent'] for j in jobs}))
        info_extra = {'계정': acct, '역할': role, '웨어하우스': wh,
                      '실행 방식': 'Snowpark 세션' if runner.snowpark else f'connector · 병렬 {workers}',
                      'Agent 기본 버전': ', '.join(f'{k}={v}' for k, v in vers.items())}

    info = dict(info_extra)
    info.update({'실행시각': dt.datetime.now().isoformat(timespec='seconds'),
                 '기준년월(자리표시)': ym, '입력 CSV': os.path.relpath(CSV_PATH, ROOT),
                 '질문 수': len(qs), '호출 수': len(jobs),
                 '부서 → Agent': '; '.join(f'{k}={"/".join(v)}' for k, v in DEPT_AGENT.items()),
                 '추가 호출(행번호)': '; '.join(f'{k}={"/".join(v)}' for k, v in EXTRA_AGENTS.items()),
                 'raw 원문': os.path.relpath(RAW_DIR, ROOT)})
    path = write_excel(jobs, results, info)
    okn = sum(1 for r in results.values() if r.get('ok'))
    print(f'\n완료 · 성공 {okn}/{len(jobs)} · Excel = {os.path.relpath(path, ROOT)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
