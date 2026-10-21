# O252 현업 전달 문서 생성 — Word 안내서 + Excel 리뷰 시트
# 수치 출처 = GN_DW.GOLD.DIM_ORG_SUCCESSION 라이브(개발계 SR61467 · 2026-10-21 build 후 조회) · tmp/o252/evidence.md
import csv
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT_DIR = '/tmp/o252_out/'
DOCX = OUT_DIR + 'O252_폐지부서_현재부서연결_안내_20261021.docx'
XLSX = OUT_DIR + 'O252_폐지부서_연결결과_검토요청_20261021.xlsx'
CSV = '/workspace/30_output_share/' + 'O252_폐지부서_귀속_현업리뷰_20261021.csv'

YEARLY = [  # 연, 전체, 현재부서 그대로, 폐지부서→현재부서 연결, 미매핑, 연결률
    ('2024', 143153, 119796, 22357, 1000, '99.3%'),
    ('2025', 153263, 133852, 18515, 896, '99.4%'),
    ('2026(진행 중)', 129707, 115856, 13230, 621, '99.5%'),
]
TOP = [  # 2024 중단건 · 부서-팀, 연결 전, 연결 후, 늘어난 건, 넘어온 폐지부서 수
    ('본부 > 회원실 > 회원서비스기획팀', 9584, 20206, 10622, 36),
    ('본부 > 회원실 > 회원참여팀', 4826, 6329, 1503, 9),
    ('국내 > 영남 > 대구사업본부', 1360, 2627, 1267, 4),
    ('국내 > 경기강원 > 경기남부사업본부', 1164, 2257, 1093, 1),
    ('국내 > 호남 > 호남권역본부', 1732, 2381, 649, 3),
    ('국내 > 서울인천 > 서울북부사업본부', 1731, 2365, 634, 2),
    ('국내 > 충청 > 대전충북사업본부', 1092, 1725, 633, 3),
    ('국내 > 영남 > 경남사업본부', 1284, 1881, 597, 2),
]

COCO_PROMPT = (
    '주간중단보고 대시보드의 부서-팀 드릴다운을 현재 조직 기준으로 바꿔줘. '
    '중단건의 실적부서 코드를 GN_DW.GOLD.DIM_ORG_SUCCESSION 테이블의 OLD_DEPT_ID 와 LEFT JOIN 하고, '
    '드릴다운 축은 CUR_ORG_PATH 를 써줘. CUR_ORG_PATH 가 비어 있으면 \'(미매핑 폐지부서)\' 로 보여줘. '
    '중단율의 분모(활동 후원 등)도 같은 방식으로 같은 테이블에 연결해서 분자와 분모의 부서 기준을 맞춰줘. '
    '바꾼 뒤 전체 합계가 바꾸기 전과 같은지 확인해줘.'
)
SQL_EXAMPLE = """SELECT COALESCE(m.CUR_ORG_PATH, '(미매핑 폐지부서)') AS 부서팀,
       COUNT(*)                AS 중단건,
       COUNT_IF(m.IS_REMAPPED) AS 폐지부서에서_넘어온_중단건
FROM GN_DW.GOLD.FACT_MEMBER_SPONSORSHIP_SPAN s
LEFT JOIN GN_DW.GOLD.DIM_ORG_SUCCESSION m
       ON m.OLD_DEPT_ID = s.SPNSR_ACMSLT_DEPT_CD
WHERE s.DSCNTC_MONTH_KEY = 202609
GROUP BY 1
ORDER BY 중단건 DESC;"""


def table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = f'{v:,}' if isinstance(v, int) else str(v)
    return t


def build_docx():
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Malgun Gothic'
    st.font.size = Pt(10)

    d.add_heading('폐지·개편 부서 → 현재 부서-팀 연결 안내', level=0)
    d.add_paragraph('대상: 주간중단보고 대시보드 및 부서별 보고서 · 작성: DW 개발 · 2026-10-21')

    d.add_heading('1. 한 줄 요약', level=1)
    d.add_paragraph(
        '없어진 부서에 잡혀 있던 중단건을 지금 살아 있는 부서-팀에 연결했습니다. '
        '2024~2026년 중단건의 99% 이상이 현재 부서-팀으로 보입니다. '
        '연결하지 못한 1% 미만은 「(미매핑 폐지부서)」로 따로 보여서 전체 합계는 바뀌지 않습니다.')

    d.add_heading('2. 바뀌는 숫자 (연도별 중단건)', level=1)
    table(d, ['연도', '전체', '원래 현재 부서', '폐지부서 → 현재 부서 연결', '미매핑', '연결률'], YEARLY)
    p = d.add_paragraph()
    r = p.add_run('※ 개발 환경 기준 숫자입니다. 운영 환경은 데이터 시점이 달라 숫자가 다를 수 있습니다(비율은 비슷합니다).')
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    d.add_heading('3. 숫자가 많이 늘어나는 부서-팀 (2024년 중단건)', level=1)
    table(d, ['부서-팀', '연결 전', '연결 후', '늘어난 건', '넘어온 폐지부서 수'], TOP)
    d.add_paragraph(
        '회원서비스기획팀이 가장 많이 늘어납니다. 없어진 회원나눔팀·e-나눔팀·회원만족팀과 폐지된 지부들의 회원이 '
        '이후 이 팀 실적으로 관리되고 있어서 그렇게 연결했습니다. 아래 5번에서 확인 부탁드립니다.')

    d.add_heading('4. 어떻게 연결했나요', level=1)
    d.add_paragraph('아래 순서로 먼저 맞는 기준을 씁니다. 각 부서가 어떤 기준으로 연결됐는지 엑셀에 적혀 있습니다.')
    for s in [
        '① 현업 확정: 현업에서 정해 준 연결이 있으면 그것을 가장 먼저 씁니다.',
        '② 같은 이름: 지금 같은 이름의 부서가 딱 하나 있으면 그 부서로 연결합니다.',
        '③ 회원 흐름: 그 부서 회원들이 이후 어느 현재 부서의 실적으로 많이 잡히는지 봅니다. '
        '(매체운영팀은 원래 실적이 많아서, 「평소보다 얼마나 더 많이 잡히는가」로 판단합니다.)',
        '④ 상위 부서: 위 기준으로 정해지지 않으면 살아 있는 상위 부서(본부·권역)로 연결합니다.',
        '⑤ 미매핑: 그래도 정할 수 없으면 「(미매핑 폐지부서)」로 둡니다.',
    ]:
        d.add_paragraph(s, style='List Bullet')
    d.add_paragraph('원래 기록된 당시 부서명은 지우지 않습니다. 필요하면 당시 부서 기준으로도 볼 수 있습니다.')

    d.add_heading('5. 확인 부탁드리는 것', level=1)
    d.add_paragraph('첨부 엑셀 「검토요청」 시트에서 아래 순서대로 봐 주세요. 맞으면 비워 두시면 됩니다.')
    for s in [
        '필수(74개 부서): 연결할 부서를 찾지 못한 곳입니다. 알고 계신 현재 부서가 있으면 적어 주세요.',
        '권장 · 상위 귀속(73개): 팀이 아니라 본부·권역 단위로 연결된 곳입니다. 더 정확한 팀이 있으면 적어 주세요.',
        '권장 · 근거 약함(36개): 회원 흐름 근거가 약한 곳입니다. 특히 지부 → 회원서비스기획팀 연결이 맞는지 봐 주세요.',
        '선택(42개): 근거가 충분한 곳입니다. 시간이 되실 때 확인해 주세요.',
    ]:
        d.add_paragraph(s, style='List Bullet')
    d.add_paragraph(
        '적는 법: 「확정_귀속부서코드」 칸에 현재 부서코드를 적습니다(엑셀 「현재부서목록」 시트 참고). '
        '한 부서가 여러 팀으로 나뉜 경우에는 대표 팀 하나를 적고 비고에 사정을 적어 주세요. '
        '연결하지 않는 게 맞으면 「연결안함」이라고 적어 주세요.')

    d.add_heading('6. 대시보드에 반영하는 방법', level=1)
    w = d.add_paragraph().add_run(
        '※ 운영 환경에 연결 테이블(GN_DW.GOLD.DIM_ORG_SUCCESSION)이 배포된 뒤에 사용해 주세요. '
        '배포가 끝나면 DW 개발에서 따로 알려 드립니다.')
    w.font.bold = True
    w.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    d.add_paragraph('CoCo 에 아래 문장을 그대로 붙여 넣으시면 됩니다.')
    q = d.add_paragraph(COCO_PROMPT)
    q.runs[0].font.bold = True
    d.add_paragraph('참고용 쿼리(2026년 9월 중단건을 현재 부서-팀별로 보기):')
    code = d.add_paragraph(SQL_EXAMPLE)
    for r in code.runs:
        r.font.name = 'Consolas'
        r.font.size = Pt(9)

    d.add_heading('7. 앞으로 관리', level=1)
    for s in [
        '현업에서 확정해 주신 연결은 다음 적재부터 자동 연결보다 우선 반영됩니다.',
        '조직 개편이 있을 때마다 연결표를 새로 정리해야 합니다. 인사 담당이나 담당자를 정해 주시면 좋겠습니다.',
        '같은 기준을 다른 부서별 보고서에도 그대로 쓸 수 있습니다(같은 테이블을 연결하면 됩니다).',
    ]:
        d.add_paragraph(s, style='List Bullet')
    d.save(DOCX)


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = '검토요청'
    head_fill = PatternFill('solid', fgColor='1F4E78')
    in_fill = PatternFill('solid', fgColor='FFF2CC')
    pri_fill = {'필수': 'F8CBAD', '권장(상위 귀속)': 'FCE4D6', '권장(근거 약함)': 'FCE4D6', '선택': 'E2EFDA'}
    with open(CSV, encoding='utf-8-sig') as fh:
        rows = list(csv.reader(fh))
    hdr = rows[0]
    order = {'필수': 0, '권장(근거 약함)': 1, '권장(상위 귀속)': 2, '선택': 3}
    pi = hdr.index('검토우선순위')
    body = sorted(rows[1:], key=lambda r: (order.get(r[pi], 9), -sum(int(r[i] or 0) for i in (5, 6, 7))))
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in body:
        vals = [int(v) if i in (5, 6, 7) and v != '' else v for i, v in enumerate(r)]
        ws.append(vals)
        row = ws.max_row
        ws.cell(row=row, column=pi + 1).fill = PatternFill('solid', fgColor=pri_fill.get(r[pi], 'FFFFFF'))
        for col in (len(hdr) - 1, len(hdr)):
            ws.cell(row=row, column=col).fill = in_fill
    widths = [12, 16, 30, 12, 14, 9, 9, 9, 18, 12, 18, 34, 12, 16, 10, 10, 9, 14, 22, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions

    g = wb.create_sheet('작성방법')
    for line in [
        '1. 「검토우선순위」가 필수 → 권장 → 선택 순으로 정렬돼 있습니다.',
        '2. 맞으면 노란 칸을 비워 두세요.',
        '3. 다른 현재 부서가 맞으면 「확정_귀속부서코드」에 코드를 적으세요(「현재부서목록」 시트 참고).',
        '4. 연결하지 않는 게 맞으면 「연결안함」이라고 적으세요.',
        '5. 한 부서가 여러 팀으로 나뉜 경우 대표 팀 하나를 적고 비고에 사정을 적어 주세요.',
        '※ 숫자는 개발 환경 기준입니다(운영 환경과 조금 다를 수 있습니다).',
    ]:
        g.append([line])
    g.column_dimensions['A'].width = 100

    import snowflake.connector
    con = snowflake.connector.connect(connection_name='default')
    try:
        cur = con.cursor()
        cur.execute(
            "SELECT DEPT_ID, DEPT_NM, ORG_PATH FROM GN_DW.GOLD.DIM_ORG_SUCCESSION s "
            "JOIN GN_DW.SILVER.CRM_ORG c ON c.DEPT_ID = s.OLD_DEPT_ID "
            "JOIN GN_DW.GOLD.DIM_ORG o ON o.ORG_SK = s.OLD_ORG_SK "
            "WHERE s.MAP_METHOD = 'SELF' ORDER BY ORG_PATH")
        act = cur.fetchall()
    finally:
        con.close()
    a = wb.create_sheet('현재부서목록')
    a.append(['현재부서코드', '현재부서명', '조직경로'])
    for c in a[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
    for r in act:
        a.append(list(r))
    a.column_dimensions['A'].width = 14
    a.column_dimensions['B'].width = 24
    a.column_dimensions['C'].width = 50
    a.freeze_panes = 'A2'
    a.auto_filter.ref = a.dimensions
    wb.save(XLSX)
    return len(body), len(act)


if __name__ == '__main__':
    build_docx()
    n, k = build_xlsx()
    print('docx', DOCX)
    print('xlsx', XLSX, 'rows', n, 'active', k)
