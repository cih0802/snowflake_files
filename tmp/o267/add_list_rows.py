"""O267 — 현업요청인수인계목록 xlsx·csv 에 O265·O266 행 3건 추가(멱등 · 이미 있으면 건너뜀)."""
import copy
import csv
import io

import openpyxl
from openpyxl.worksheet.cell_range import CellRange

BASE = "/workspace/98_현업요청인수인계/현업요청인수인계목록"

ROWS = [
    {
        "no": "13",
        "scope": "현업 요청 대응 · 알림톡 열람·발송구분",
        "main": "김나경(회원서비스기획팀) · 알림톡 열람 = GA 링크 클릭 · 알림톡/문자/이메일 발송구분(대·중·소) 질문 · 열람 2축(A·B)",
        "id": "O265-1",
        "attach": "열람 정의(GA 링크 클릭 · 회원번호 파라미터 연결), 열람 A(그 템플릿 링크 클릭)·B(기간 내 어떤 알림톡이든 클릭) 판정 원리, 세 채널 발송구분(대·중·소) 배선, 질문 키워드 안내, 서비스(명)·서비스(건) 값 증가 공지(이메일 유입).",
        "check": "회신문 내용대로 Agent 에 열람 A·B 와 발송구분(대·중·소) 질문을 해 보고 답변 기준이 회신과 같은지 확인합니다.",
        "file": "`20261028_김나경_알림톡열람_발송구분_회신.md` v1.0 (2026-10-28)",
        "loc": "`98_현업요청인수인계/20261028_김나경_알림톡열람_발송구분_회신.md` §1~§6",
    },
    {
        "no": "14",
        "scope": "현업 요청 대응 · 발송 이력·알림톡 클릭 판정",
        "main": "강지영(회원참여팀) · 신규/증액/재후원 발송 이력 연결 기간 · 알림톡 대·중·소 배선 · 특정 알림톡 클릭 판정 3방식",
        "id": "O265-2",
        "attach": "발송 이력 = 사건 전후 기간 연결(당일만 아님), 알림톡 대·중·소 배선, 특정 알림톡 클릭 판정 3방식과 우선순위(질문 시 템플릿 정보 > 고정 매핑표 > 추정표), 고정 매핑표 선행조건(IT 웹팀 리다이렉트 대상표), 단축·리다이렉트 링크로 추정이 필요한 이유, 서비스(명)·서비스(건) 값 증가 공지.",
        "check": "템플릿 정보 유무별로 같은 질문을 해 판정 방식이 회신의 우선순위대로 바뀌는지 확인합니다. 고정 매핑표는 IT(웹팀) 대상표 수령 여부를 확인합니다.",
        "file": "`20261028_강지영_발송이력_알림톡클릭_회신.md` v1.0 (2026-10-28)",
        "loc": "`98_현업요청인수인계/20261028_강지영_발송이력_알림톡클릭_회신.md` §1~§7",
    },
    {
        "no": "15",
        "scope": "현업 요청 대응 · 마케팅 Agent 피드백",
        "main": "마케팅 피드백 1~4 · 디지털 최종정산금액 · 캠페인별 광고비/개발건수/제작비 한 표 · 캠페인명·채널 조회 · LTV 답변 항목 고정",
        "id": "O266-1",
        "attach": "1) 디지털 광고비 문의 시 최종정산금액 병기, 2) 마케팅캠페인별 광고비·개발건수·제작비 결합 표(제작비 세세목 정의 요청), 3) 캠페인명/채널 조회, 4) LTV 답변 항목 고정 및 유지개월(무기한) 계산 기준 변경.",
        "check": "피드백 4문항을 Agent 에 다시 질문해 회신문 항목대로 답하는지 확인합니다. 제작비 세세목 목록을 회신합니다.",
        "file": "`20261028_마케팅Agent_피드백_회신.md` v1.0 (2026-10-28)",
        "loc": "`98_현업요청인수인계/20261028_마케팅Agent_피드백_회신.md` 전체(「무엇이 바뀌었나」~「함께 반영된 사항」)",
    },
]


def do_xlsx():
    path = BASE + ".xlsx"
    wb = openpyxl.load_workbook(path)
    # Sheet2(상세) — 표 끝에 덧붙인다
    ws2 = wb["현업요청인수인계목록"]
    have2 = {str(ws2.cell(r, 4).value) for r in range(1, ws2.max_row + 1)}
    last2 = max(r for r in range(1, ws2.max_row + 1) if ws2.cell(r, 1).value)
    for row in ROWS:
        if row["id"] in have2:
            continue
        last2 += 1
        vals = [row["no"], row["scope"], row["main"], row["id"], row["scope"],
                row["attach"], row["check"], row["file"].replace("`", ""), row["loc"].replace("`", "")]
        for c, v in enumerate(vals, 1):
            src = ws2.cell(last2 - 1, c)
            dst = ws2.cell(last2, c, v)
            dst._style = copy.copy(src._style)
    # Sheet1(확인서) — 12번 행 뒤에 끼우고 아래 병합을 옮긴다
    ws1 = wb["Sheet1"]
    have1 = {str(ws1.cell(r, 4).value) for r in range(1, ws1.max_row + 1)}
    r12 = next(r for r in range(1, ws1.max_row + 1) if str(ws1.cell(r, 4).value) == "D12")
    new = [row for row in ROWS if row["id"] not in have1]
    if new:
        n = len(new)
        below = [CellRange(str(m)) for m in ws1.merged_cells.ranges if m.min_row > r12]
        for m in below:
            ws1.merged_cells.remove(m) if hasattr(ws1.merged_cells, "remove") else ws1.unmerge_cells(str(m))
        heights = {r: ws1.row_dimensions[r].height for r in range(r12 + 1, ws1.max_row + 1)}
        ws1.insert_rows(r12 + 1, n)
        for r, h in heights.items():
            ws1.row_dimensions[r + n].height = h
        for m in below:
            ws1.merge_cells(start_row=m.min_row + n, start_column=m.min_col,
                            end_row=m.max_row + n, end_column=m.max_col)
        for i, row in enumerate(new):
            r = r12 + 1 + i
            vals = [row["no"], row["scope"], row["main"], row["id"], None, None, "미확인", None]
            for c, v in enumerate(vals, 1):
                dst = ws1.cell(r, c, v)
                dst._style = copy.copy(ws1.cell(r12, c)._style)
            ws1.row_dimensions[r].height = ws1.row_dimensions[r12].height
    import shutil
    wb.save("/tmp/o267_list.xlsx")
    shutil.copyfile("/tmp/o267_list.xlsx", path)


def do_csv():
    path = BASE + ".csv"
    raw = open(path, "rb").read().decode("cp949")
    rows = list(csv.reader(io.StringIO(raw)))
    have = {r[3] for r in rows if len(r) > 3}
    for row in ROWS:
        if row["id"] in have:
            continue
        rows.append([row["no"], row["scope"], row["main"], row["id"], row["scope"],
                     row["attach"], row["check"], row["file"].replace("`", ""), row["loc"].replace("`", "")])
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(rows)
    open(path, "wb").write(buf.getvalue().encode("cp949"))


if __name__ == "__main__":
    do_xlsx()
    do_csv()
    print("OK")
