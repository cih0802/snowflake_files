# 탭 4 월별목표 참조테이블: 엑셀 '2026 목표_회원개발(월별)' 255~269행 → MONTHLY_TARGET_REF (전체 교체)
import io
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime

import pandas as pd
import streamlit as st

REF_TABLE = "GN_DW.LAB.MONTHLY_TARGET_REF"
BAK_TABLE = "GN_DW.LAB.MONTHLY_TARGET_REF_BAK"

DEFAULT_SHEET = "2026 목표_회원개발(월별)"
DEFAULT_HEADER_ROW = 255
DEFAULT_FIRST_ROW = 256
DEFAULT_LAST_ROW = 269
HEADER_TITLES = ["사업장명", "교육기관/지역사회분류", "캠페인분류"]
KNOWN_P_EXCEPTIONS = {"지역사회(수기)"}  # 268행 P열 100,000은 엑셀 입력 오류로 확인됨(2026-10-28) → 월목표 0 적재, P 무시

# 엑셀 C열 라벨 → (B열 분류, 대시보드 행키)
EXPECTED = [
    ("가족그림", "교육", "E_FAMILY"),
    ("유아기타", "교육", "E_KID_ETC"),
    ("희망편지", "교육", "E_HOPE"),
    ("초등기타", "교육", "E_ELEM_ETC"),
    ("청소년 캠페인", "교육", "E_TEEN"),
    ("교육기관(기타)", "교육", "E_ETC"),
    ("신규기획사업", "지역", "R_NEW"),
    ("교회 캠페인", "지역", "R_CHURCH"),
    ("대학생 캠페인", "지역", "R_UNIV"),
    ("좋은이웃가게", "지역", "R_SHOP"),
    ("가두 캠페인", "지역", "R_STREET"),
    ("지역사회(기타)", "지역", "R_ETC"),
    ("지역사회(수기)", "지역", "R_MANUAL"),
    ("기업오프라인", "", "HQ"),
]
LABEL_TO_KEY = {lab: key for lab, _, key in EXPECTED}
ANNUAL_SADAN_LABEL = "가족그림"  # 첫 데이터 행 Q열 = 사단 연목표
ANNUAL_TOTAL_LABEL = "기업오프라인"  # 마지막 데이터 행 Q열 = 총계 연목표

MONTH_COLS = [chr(ord("D") + i) for i in range(12)]  # D~O
CUM_COLS = ["S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC", "AD"]  # 1~12월 누적 (AE열은 대조하지 않음)


# ─── 엑셀 읽기 (openpyxl 있으면 사용, 없으면 표준 라이브러리로 저장된 값 읽기) ───
def _col_to_idx(col):
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def _idx_to_col(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _read_openpyxl(data, sheet, r1, r2):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"시트 '{sheet}' 없음. 시트 목록: {wb.sheetnames}")
    ws = wb[sheet]
    cells = {}
    for r_idx, row in enumerate(ws.iter_rows(min_row=r1, max_row=r2, max_col=31, values_only=True), start=r1):
        for c_idx, v in enumerate(row, start=1):
            if v is not None:
                cells[(r_idx, _idx_to_col(c_idx))] = v
    return cells


def _read_stdlib(data, sheet, r1, r2):
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    z = zipfile.ZipFile(io.BytesIO(data))
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {r.get("Id"): r.get("Target") for r in rels}
    sheets = {s.get("name"): s.get(f"{{{ns['r']}}}id") for s in wb.find("m:sheets", ns)}
    if sheet not in sheets:
        raise ValueError(f"시트 '{sheet}' 없음. 시트 목록: {list(sheets)}")
    target = rid_to_target[sheets[sheet]].lstrip("/")
    path = target if target.startswith("xl/") else f"xl/{target}"
    shared = []
    if "xl/sharedStrings.xml" in z.namelist():
        for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall("m:si", ns):
            shared.append("".join(t.text or "" for t in si.iter(f"{{{ns['m']}}}t")))
    cells = {}
    for row in ET.fromstring(z.read(path)).iter(f"{{{ns['m']}}}row"):
        r = int(row.get("r"))
        if r < r1 or r > r2:
            continue
        for c in row.findall("m:c", ns):
            col = re.match(r"[A-Z]+", c.get("r")).group(0)
            t, v = c.get("t"), c.find("m:v", ns)
            if t == "inlineStr":
                val = "".join(x.text or "" for x in c.iter(f"{{{ns['m']}}}t"))
            elif v is None or v.text is None:
                continue
            elif t == "s":
                val = shared[int(v.text)]
            elif t in ("str", "e"):
                val = v.text
            elif t == "b":
                val = v.text == "1"
            else:
                f = float(v.text)
                val = int(f) if f.is_integer() else f
            cells[(r, col)] = val
    return cells


def read_cells(data, sheet, r1, r2):
    try:
        return _read_openpyxl(data, sheet, r1, r2), "openpyxl"
    except ImportError:
        return _read_stdlib(data, sheet, r1, r2), "표준 라이브러리(zip/xml)"


# ─── 검증 ───
def _txt(v):
    return "" if v is None else str(v).strip()


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def parse_and_validate(cells, header_row, first_row, last_row):
    """반환: dict(rows, errors, confirm, p_diffs, cum_diffs, header_found)"""
    errors, confirm, p_diffs, cum_diffs = [], [], [], []

    # 헤더 위치 탐색 (지정 행 ±20 범위)
    header_found = [r for r in range(header_row - 20, header_row + 21)
                    if [_txt(cells.get((r, c))) for c in "ABC"] == HEADER_TITLES]
    hdr = [_txt(cells.get((header_row, c))) for c in "ABC"]
    months = [_txt(cells.get((header_row, c))) for c in MONTH_COLS]
    month_ok = all(re.fullmatch(rf"0?{i + 1}\s*월", m) for i, m in enumerate(months))
    if hdr != HEADER_TITLES or not month_ok:
        msg = f"{header_row}행이 헤더가 아님 (A~C={hdr}, D~O={months})."
        if header_found:
            msg += f" 헤더가 {header_found}행에서 발견됨 → 헤더행/시작행/끝행 값을 조정하세요."
        errors.append(msg)

    n_rows = last_row - first_row + 1
    if n_rows != len(EXPECTED):
        errors.append(f"데이터 행 수 {n_rows}개 (기대 {len(EXPECTED)}개: {first_row}~{last_row}행)")

    rows = []
    for i, r in enumerate(range(first_row, last_row + 1)):
        label, grp = _txt(cells.get((r, "C"))), _txt(cells.get((r, "B")))
        exp = EXPECTED[i] if i < len(EXPECTED) else None
        if exp is None or label != exp[0] or grp != exp[1]:
            errors.append(f"{r}행 라벨 불일치: 엑셀 B='{grp}', C='{label}' / 기대 "
                          + (f"B='{exp[1]}', C='{exp[0]}'" if exp else "(행 없음)"))
        vals = []
        for mi, col in enumerate(MONTH_COLS):
            raw = cells.get((r, col))
            v = _num(raw)
            if v is None:
                confirm.append(f"{r}행 {label} {mi + 1}월({col}{r}) 값 = {raw!r} (숫자 아님/빈 셀)")
            vals.append(v)
        p, q = _num(cells.get((r, "P"))), _num(cells.get((r, "Q")))
        s = sum(v for v in vals if v is not None)
        if p is None or abs(s - p) > 0.5:
            p_diffs.append({"행": r, "라벨": label, "D~O 합": s, "P 합계": p,
                            "구분": "알려진 예외(P열 입력 오류 확인됨, 월목표 0 적재)" if label in KNOWN_P_EXCEPTIONS else "중단 사유"})
        cum = 0.0
        for mi, col in enumerate(CUM_COLS):
            cum += vals[mi] or 0.0
            xc = _num(cells.get((r, col)))
            if xc is None or abs(xc - cum) > 0.5:
                cum_diffs.append({"행": r, "라벨": label, "월": mi + 1, "SQL 누계": cum, f"엑셀 누적({col})": xc})
        rows.append({"SRC_ROW": r, "CATEGORY_GROUP": grp or None, "CATEGORY": label,
                     "MONTHS": vals, "P": p, "ANNUAL_TARGET": q})

    for d in p_diffs:
        if d["구분"] == "중단 사유":
            errors.append(f"{d['행']}행 {d['라벨']}: D~O 합 {d['D~O 합']:,.0f} ≠ P {d['P 합계']}")
    return {"rows": rows, "errors": errors, "confirm": confirm, "p_diffs": p_diffs,
            "cum_diffs": cum_diffs, "header_found": header_found}


# ─── 적재 (전체 교체, 단일 트랜잭션) ───
def refresh_table(conn, rows, year, src_file):
    loaded_at = datetime.now().replace(microsecond=0)
    recs = []
    for r in rows:
        for mi, v in enumerate(r["MONTHS"]):
            recs.append((year, mi + 1, r["CATEGORY_GROUP"], r["CATEGORY"], v,
                         r["ANNUAL_TARGET"], r["SRC_ROW"], loaded_at, src_file))
    raw = conn.raw_connection
    # 연결의 바인딩 방식을 따름 (앱의 conn.query는 qmark '?' 사용)
    style = getattr(raw, "_paramstyle", None) or "pyformat"
    ph = ", ".join(["?"] * 9) if style == "qmark" else ", ".join(["%s"] * 9)
    cur = raw.cursor()
    try:
        cur.execute("BEGIN")
        cur.execute(f"INSERT OVERWRITE INTO {BAK_TABLE} SELECT * FROM {REF_TABLE}")
        cur.execute(f"DELETE FROM {REF_TABLE}")
        cur.executemany(
            f"INSERT INTO {REF_TABLE} (TARGET_YEAR, MONTH_NO, CATEGORY_GROUP, CATEGORY, TARGET_VALUE, "
            f"ANNUAL_TARGET, SRC_ROW, LOADED_AT, SRC_FILE) VALUES ({ph})",
            recs,
        )
        cur.execute("COMMIT")
    except Exception:
        cur.execute("ROLLBACK")
        raise
    finally:
        cur.close()
    return len(recs), loaded_at


REF_CSV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "monthly_target_ref.csv")


@st.cache_data(ttl=600)
def load_ref(_conn, year):
    """참조테이블 조회 → 없으면 앱 폴더 CSV. 둘 다 없거나 해당 연도 없음이면 (None, 사유)"""
    reason = None
    try:
        df = _conn.query(
            f"SELECT TARGET_YEAR, MONTH_NO, CATEGORY_GROUP, CATEGORY, TARGET_VALUE, ANNUAL_TARGET, SRC_ROW, "
            f"LOADED_AT, SRC_FILE FROM {REF_TABLE} WHERE TARGET_YEAR = ?",
            params=[year], ttl=0,
        )
    except Exception as e:  # noqa: BLE001 - 테이블 미생성 시 CSV로 대체
        reason = f"{REF_TABLE} 조회 불가: {str(e).splitlines()[0][:120]}"
        df = None
    if df is None or df.empty:
        if not os.path.exists(REF_CSV):
            return None, reason or f"{REF_TABLE}에 {year}년 목표 없음"
        df = pd.read_csv(REF_CSV, dtype={"CATEGORY_GROUP": str})
        df = df[df["TARGET_YEAR"] == year]
        df["LOADED_AT"] = datetime.fromtimestamp(os.path.getmtime(REF_CSV)).replace(microsecond=0)
        if df.empty:
            return None, f"{year}년 목표 없음 (테이블·CSV)"
    if "SITE" not in df.columns:  # 테이블에는 사업장 행이 없음 → 캠페인 목표만
        df["SITE"] = ""
    df["SITE"] = df["SITE"].fillna("").astype(str)
    df["MONTH_NO"] = df["MONTH_NO"].astype(int)
    df["TARGET_VALUE"] = df["TARGET_VALUE"].astype(float)
    df["ANNUAL_TARGET"] = pd.to_numeric(df["ANNUAL_TARGET"], errors="coerce")
    df["ROWKEY"] = df["CATEGORY"].where(df["SITE"] == "").map(LABEL_TO_KEY)
    return df, None


# ─── 갱신 화면 ───
def render_refresh_ui(conn, default_year):
    st.markdown("##### 월별목표 참조테이블 갱신")
    st.caption(f"대상: {REF_TABLE} (직전 버전 → {BAK_TABLE}) · 엑셀 업로드 → 미리보기·검증 → 갱신")
    up = st.file_uploader("월별목표 엑셀 (.xlsx)", type=["xlsx"], key="t4ref_file")
    c1, c2, c3, c4, c5 = st.columns([3, 1, 1, 1, 1])
    sheet = c1.text_input("시트명", DEFAULT_SHEET, key="t4ref_sheet")
    header_row = c2.number_input("헤더행", value=DEFAULT_HEADER_ROW, step=1, key="t4ref_hdr")
    first_row = c3.number_input("시작행", value=DEFAULT_FIRST_ROW, step=1, key="t4ref_r1")
    last_row = c4.number_input("끝행", value=DEFAULT_LAST_ROW, step=1, key="t4ref_r2")
    year = c5.number_input("대상연도", value=int(default_year), step=1, key="t4ref_year")
    if up is None:
        return

    try:
        cells, engine = read_cells(up.getvalue(), sheet, int(header_row) - 20, max(int(last_row), int(header_row) + 20))
    except Exception as e:  # noqa: BLE001
        st.error(f"엑셀 읽기 실패: {e}")
        return
    res = parse_and_validate(cells, int(header_row), int(first_row), int(last_row))
    st.caption(f"읽기 방식: {engine} (수식은 마지막 저장된 계산값 사용)")

    prev = pd.DataFrame(
        [[r["SRC_ROW"], r["CATEGORY_GROUP"] or "", r["CATEGORY"]] + r["MONTHS"] + [sum(v or 0 for v in r["MONTHS"]), r["P"], r["ANNUAL_TARGET"]]
         for r in res["rows"]],
        columns=["행", "분류", "캠페인분류"] + [f"{m}월" for m in range(1, 13)] + ["D~O 합", "P 합계", "Q 연목표"],
    )
    st.dataframe(prev, hide_index=True, use_container_width=True)

    for e in res["errors"]:
        st.error(e)
    if res["p_diffs"]:
        st.warning("행별 D~O 합 ≠ P 합계")
        st.dataframe(pd.DataFrame(res["p_diffs"]), hide_index=True, use_container_width=True)
    if res["cum_diffs"]:
        st.warning(f"엑셀 누적(S~AD) ≠ SQL 계산 누계: {len(res['cum_diffs'])}건")
        st.dataframe(pd.DataFrame(res["cum_diffs"]), hide_index=True, use_container_width=True)
    ok_confirm = True
    if res["confirm"]:
        st.warning("숫자가 아니거나 빈 셀 (0으로 채우지 않음 → NULL로 적재)")
        st.write("\n".join(f"- {x}" for x in res["confirm"]))
        ok_confirm = st.checkbox("위 셀을 NULL(목표 없음)로 적재하는 데 동의합니다", key="t4ref_confirm")

    can_load = not res["errors"] and ok_confirm
    if st.button("참조테이블 갱신", type="primary", disabled=not can_load, key="t4ref_go"):
        try:
            n_rows, loaded_at = refresh_table(conn, res["rows"], int(year), up.name)
        except Exception as e:  # noqa: BLE001
            st.error(f"갱신 실패 (기존 값 유지): {e}")
            return
        load_ref.clear()
        chk = conn.query(
            f"SELECT MONTH_NO, SUM(TARGET_VALUE) AS 월합계, COUNT(*) AS 행수 FROM {REF_TABLE} "
            f"WHERE TARGET_YEAR = ? GROUP BY 1 ORDER BY 1", params=[int(year)], ttl=0,
        )
        st.success(f"갱신 완료: {n_rows}행 적재 · 적재 시각 {loaded_at:%Y-%m-%d %H:%M:%S} · 파일 {up.name}")
        st.dataframe(chk, hide_index=True, use_container_width=True)
