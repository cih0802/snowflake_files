#!/usr/bin/env python3
"""트랙 B 구조 산출물 생성 — _ENV/<계정>/B1 csv + dbt 모델·sources + Agent spec → D04 xlsx/csv · D05 FK csv+mermaid md · D06 테이블 계보 csv · D08 SV 요약 csv · D09 Agent 도구 csv

🆕 [UNLABELED 신설 · `new_tool.py` 경유 ⇒ `gate_census` 등재 완료]
🔴 분류 = `GEN`. 이 분류의 계약을 지켜라(정본 = `scripts/gate_census.py`).
🔴 산출물은 자동 생성물이다 — 손으로 고치면 다음 생성에서 사라진다.

사용법
    python3 scripts/b_struct_build.py <계정>        # 예: IW20760
    선행 = python3 scripts/env_capture.py --id B1 "85_실제 인수인계 문서/_근거쿼리/B1_구조_근거.sql"

입력
    · _ENV/<계정>/B1_Q01~Q11 csv (라이브 구조 · [공통])
    · 10_dbt_pipeline/models/**/*.sql · models/silver/_sources.yml (dbt 계보 · 파일)
    · cortex_project/agents/*/agent_spec.yaml (Agent 도구 · 파일)
출력(85_실제 인수인계 문서/20_정식산출물/ · 첫 줄 각주에 계정·생성 시각)
    · D04_테이블컬럼정의서.xlsx  시트 = 요약 · 테이블목록 · 레이어별 컬럼(BRONZE/SILVER/GOLD/SERVING/ML/MSTR/OPS)
    · D04_테이블목록.csv · D04_컬럼정의.csv
    · D05_FK관계.csv · D05_ERD_GOLD_자동생성.md (mermaid erDiagram)
    · D06_테이블계보.csv  (dbt ref·source + SV 논리테이블→base + 뷰 의존 · 테이블 단위)
    · D08_SV요약.csv · D08_SV메트릭.csv
    · D09_Agent도구.csv
    · D01_지표매핑.csv  (선택 · _ENV/<계정>/B2_Q01_1.csv 가 있을 때 · 지표 → 도구 · SV 필드 · 대응 상태)

종료코드 = 0 성공 · 1 입력 부재 · 2 사용법 오류.
"""
import collections
import csv
import datetime
import glob
import io
import os
import re
import sys

BASE = '85_실제 인수인계 문서'
OUT = os.path.join(BASE, '20_정식산출물')
DBT = '10_dbt_pipeline'
AGENTS = 'cortex_project/agents'
LAYERS = ['BRONZE', 'SILVER', 'GOLD', 'SERVING', 'ML', 'MSTR', 'OPS']


def rd(path):
    lines = [l for l in io.open(path, encoding='utf-8').read().splitlines() if not l.startswith('#')]
    return list(csv.DictReader(lines))


def layer(schema):
    for L in LAYERS:
        if schema.startswith(L):
            return L
    return 'ETC'


def write_csv(name, head, rows, foot):
    sio = io.StringIO()
    sio.write(foot + '\n')
    w = csv.writer(sio, lineterminator='\n')
    w.writerow(head)
    w.writerows(rows)
    io.open(os.path.join(OUT, name), 'w', encoding='utf-8').write(sio.getvalue())


def cut(s, n=300):
    s = (s or '').replace('\r', '').replace('\n', ' ⏎ ')
    return s if len(s) <= n else s[:n] + '…'


def dbt_lineage():
    """dbt 모델 파일명 = 물리 테이블명. ref·source 를 테이블 단위 간선으로."""
    import yaml
    src = {}
    for f in glob.glob(os.path.join(DBT, 'models', '**', '*.yml'), recursive=True):
        d = yaml.safe_load(open(f, encoding='utf-8')) or {}
        for s in d.get('sources', []) or []:
            sch = str(s.get('schema') or '')
            if '{{' in sch:  # bronze_crm 의 dev2 분기 → 기본 target 값
                m = re.search(r"else\s+'([A-Z_]+)'", sch)
                sch = m.group(1) if m else sch
            for t in s.get('tables', []) or []:
                src[(s['name'], t['name'])] = (sch.upper(), t['name'].upper())
    models, edges = {}, []
    for f in glob.glob(os.path.join(DBT, 'models', '**', '*.sql'), recursive=True):
        name = os.path.splitext(os.path.basename(f))[0].upper()
        parts = f.split(os.sep)
        lyr = parts[parts.index('models') + 1].upper() if 'models' in parts else ''
        models[name] = (lyr, f)
        t = open(f, encoding='utf-8', errors='replace').read()
        for m in re.finditer(r"ref\(\s*'([^']+)'\s*\)", t):
            edges.append((name, 'ref', m.group(1).upper()))
        for m in re.finditer(r"source\(\s*'([^']+)'\s*,\s*'([^']+)'\s*\)", t):
            sch, tbl = src.get((m.group(1), m.group(2)), ('?' + m.group(1), m.group(2).upper()))
            edges.append((name, 'source', sch + '.' + tbl))
    return models, sorted(set(edges))


def agent_tools():
    import yaml
    rows = []
    for f in sorted(glob.glob(os.path.join(AGENTS, '*', 'agent_spec.yaml'))):
        a = f.split(os.sep)[-2]
        d = yaml.safe_load(open(f, encoding='utf-8'))
        res = d.get('tool_resources', {}) or {}
        for t in d.get('tools', []) or []:
            ts = t.get('tool_spec', {})
            r = res.get(ts.get('name'), {}) or {}
            target = r.get('semantic_view') or r.get('search_service') or ''
            rows.append([a, ts.get('name'), ts.get('type'), target.split('.')[-1], cut(ts.get('description'), 200)])
    return rows


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    acct = sys.argv[1].upper()
    env = os.path.join(BASE, '_ENV', acct)
    need = ['B1_Q01_1', 'B1_Q02_1', 'B1_Q03_1', 'B1_Q03_2', 'B1_Q04_1', 'B1_Q05_1', 'B1_Q06_1', 'B1_Q08_1', 'B1_Q09_1', 'B1_Q11_1']
    miss = [n for n in need if not os.path.exists(os.path.join(env, n + '.csv'))]
    if miss:
        print('🔴 입력 부재: %s' % miss)
        return 1
    os.makedirs(OUT, exist_ok=True)
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    foot = '# 자동 생성 = scripts/b_struct_build.py · account: %s · 생성 %s · 원천 = _ENV/%s/B1_*.csv · 손으로 고치지 않는다' % (acct, now, acct)

    tables = rd(os.path.join(env, 'B1_Q01_1.csv'))
    cols = rd(os.path.join(env, 'B1_Q02_1.csv'))
    pks = rd(os.path.join(env, 'B1_Q03_1.csv'))
    fks = rd(os.path.join(env, 'B1_Q03_2.csv'))
    svs = rd(os.path.join(env, 'B1_Q04_1.csv'))
    svt = rd(os.path.join(env, 'B1_Q05_1.csv'))
    mets = rd(os.path.join(env, 'B1_Q06_1.csv'))
    deps = rd(os.path.join(env, 'B1_Q09_1.csv'))
    vqr = {r['SV_NAME']: r['N_VQR'] for r in rd(os.path.join(env, 'B1_Q11_1.csv'))}

    pk_set = {(r['schema_name'], r['table_name'], r['column_name']) for r in pks}
    fk_map = {(r['fk_schema_name'], r['fk_table_name'], r['fk_column_name']):
              '%s.%s.%s' % (r['pk_schema_name'], r['pk_table_name'], r['pk_column_name']) for r in fks}
    ncol = collections.Counter((c['TABLE_SCHEMA'], c['TABLE_NAME']) for c in cols)
    ncmt = collections.Counter((c['TABLE_SCHEMA'], c['TABLE_NAME']) for c in cols if (c['COMMENT'] or '').strip())

    # ---- D04 ----
    H_T = ['LAYER', 'SCHEMA', 'TABLE', 'TYPE', 'N_COLUMNS', 'N_COL_COMMENT', 'ROW_COUNT', 'BYTES', 'OWNER', 'COMMENT']
    trows = [[layer(t['TABLE_SCHEMA']), t['TABLE_SCHEMA'], t['TABLE_NAME'], t['TABLE_TYPE'],
              ncol[(t['TABLE_SCHEMA'], t['TABLE_NAME'])], ncmt[(t['TABLE_SCHEMA'], t['TABLE_NAME'])],
              t['ROW_COUNT'], t['BYTES'], t['TABLE_OWNER'], cut(t['COMMENT'])] for t in tables]
    H_C = ['LAYER', 'SCHEMA', 'TABLE', 'POS', 'COLUMN', 'DATA_TYPE', 'LENGTH', 'PRECISION', 'SCALE', 'NULLABLE',
           'PK', 'FK_REF', 'COMMENT']
    crows = []
    for c in cols:
        k = (c['TABLE_SCHEMA'], c['TABLE_NAME'], c['COLUMN_NAME'])
        crows.append([layer(c['TABLE_SCHEMA']), c['TABLE_SCHEMA'], c['TABLE_NAME'], c['ORDINAL_POSITION'], c['COLUMN_NAME'],
                      c['DATA_TYPE'], c['CHARACTER_MAXIMUM_LENGTH'], c['NUMERIC_PRECISION'], c['NUMERIC_SCALE'],
                      c['IS_NULLABLE'], 'Y' if k in pk_set else '', fk_map.get(k, ''), cut(c['COMMENT'])])
    write_csv('D04_테이블목록.csv', H_T, trows, foot)
    write_csv('D04_컬럼정의.csv', H_C, crows, foot)

    summ = collections.defaultdict(lambda: [0, 0, 0, 0])
    for r in trows:
        s = summ[(r[0], r[1])]
        s[0] += r[3] == 'BASE TABLE'
        s[1] += r[3] == 'VIEW'
        s[2] += r[4]
        s[3] += r[5]
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = '요약'
    ws.append([foot.lstrip('# ')])
    ws.append(['LAYER', 'SCHEMA', 'TABLES', 'VIEWS', 'COLUMNS', 'COLUMNS_WITH_COMMENT', 'COMMENT_RATE(%)'])
    for (L, S), v in sorted(summ.items()):
        ws.append([L, S, v[0], v[1], v[2], v[3], round(100.0 * v[3] / v[2], 2) if v[2] else None])
    s = wb.create_sheet('테이블목록')
    s.append(H_T)
    for r in trows:
        s.append(r)
    for L in LAYERS:
        s = wb.create_sheet(L)
        s.append(H_C)
        for r in crows:
            if r[0] == L:
                s.append(r)
    for sh in wb.worksheets:
        for c in sh[2 if sh.title == '요약' else 1]:
            c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)  # 🔴 마운트 seek 미지원 ⇒ 메모리 저장 후 1회 쓰기(a2_matrix_build 와 같은 이유)
    with open(os.path.join(OUT, 'D04_테이블컬럼정의서.xlsx'), 'wb') as fh:
        fh.write(buf.getvalue())

    # ---- D05 ----
    fk_rows = sorted({(r['fk_name'], r['fk_schema_name'], r['fk_table_name'], r['fk_column_name'],
                       r['pk_schema_name'], r['pk_table_name'], r['pk_column_name']) for r in fks})
    write_csv('D05_FK관계.csv', ['FK_NAME', 'FK_SCHEMA', 'FK_TABLE', 'FK_COLUMN', 'PK_SCHEMA', 'PK_TABLE', 'PK_COLUMN'],
              fk_rows, foot)
    gold_t = sorted({(r[2]) for r in fk_rows if r[1] == 'GOLD'} | {r[5] for r in fk_rows if r[4] == 'GOLD'})
    lines = ['<!-- %s -->' % foot.lstrip('# '), '', '# D05 부록 — GOLD 물리 ERD (FK 실측 · 자동 생성)', '',
             '- 원천 = `SHOW IMPORTED KEYS IN DATABASE GN_DW`(B1 Q03-2) · FK %d건 · GOLD 테이블 %d개' % (len(fk_rows), len(gold_t)),
             '- 관계선 라벨 = FK 컬럼명 · 설명은 `D05_ERD.md` 참조', '', '```mermaid', 'erDiagram']
    for r in fk_rows:
        if r[1] == 'GOLD' and r[4] == 'GOLD':
            lines.append('  %s ||--o{ %s : "%s"' % (r[5], r[2], r[3]))
    lines += ['```', '']
    io.open(os.path.join(OUT, 'D05_ERD_GOLD_자동생성.md'), 'w', encoding='utf-8').write('\n'.join(lines))

    # ---- D06 ----
    models, edges = dbt_lineage()
    phys = {(t['TABLE_SCHEMA'], t['TABLE_NAME']) for t in tables}
    by_name = collections.defaultdict(list)
    for (S, T) in phys:
        by_name[T].append(S)
    lin = []

    def where(name):
        """dbt 모델이면 그 레이어 스키마(SILVER/GOLD) — 같은 이름이 여러 스키마에 있어도 모호하지 않게. 아니면 물리 스키마들."""
        if name in models and models[name][0] in ('SILVER', 'GOLD') and models[name][0] in by_name.get(name, []):
            return models[name][0]
        return '/'.join(sorted(by_name.get(name, []))) or '?'

    for (m, kind, up) in edges:
        down_s = where(m)
        if kind == 'ref':
            lin.append(['dbt ref', where(up) + '.' + up, down_s + '.' + m, models.get(m, ('', ''))[1]])
        else:
            lin.append(['dbt source', up, down_s + '.' + m, models.get(m, ('', ''))[1]])
    for r in svt:
        lin.append(['SV base', r['BASE_TABLE_SCHEMA'] + '.' + r['BASE_TABLE_NAME'],
                    'SERVING.' + r['SEMANTIC_VIEW_NAME'] + ' (' + r['LOGICAL_TABLE'] + ')', 'INFORMATION_SCHEMA.SEMANTIC_TABLES'])
    for r in deps:
        if r['REFERENCING_OBJECT_DOMAIN'] in ('VIEW', 'MATERIALIZED VIEW'):
            lin.append(['view dependency', r['REFERENCED_SCHEMA'] + '.' + r['REFERENCED_OBJECT_NAME'],
                        r['REFERENCING_SCHEMA'] + '.' + r['REFERENCING_OBJECT_NAME'], 'ACCOUNT_USAGE.OBJECT_DEPENDENCIES'])
    lin = sorted({tuple(x) for x in lin})
    write_csv('D06_테이블계보.csv', ['EDGE_TYPE', 'UPSTREAM', 'DOWNSTREAM', 'EVIDENCE'], lin, foot)
    broken = [x for x in lin if x[1].startswith('?') or x[2].startswith('?') or x[1].startswith('?.') or '?.' in x[1]]
    nophys = sorted(m for m in models if not by_name.get(m))

    # ---- D08 ----
    wire = collections.defaultdict(list)
    tools = agent_tools()
    for a, n, t, tgt, _ in tools:
        if tgt.startswith('SV_'):
            wire[tgt].append(a.replace('AGENT_', ''))
    sv_rows = [[r['SV_NAME'], r['BASE_SCHEMAS'], r['N_TABLES'], r['N_DIMENSIONS'], r['N_FACTS'], r['N_METRICS'],
                r['N_RELATIONSHIPS'], vqr.get(r['SV_NAME'], ''), '+'.join(sorted(wire.get(r['SV_NAME'], [])))] for r in svs]
    write_csv('D08_SV요약.csv', ['SV_NAME', 'BASE_SCHEMAS', 'N_TABLES', 'N_DIMENSIONS', 'N_FACTS', 'N_METRICS',
                                 'N_RELATIONSHIPS', 'N_VQR', 'AGENTS'], sv_rows, foot)
    write_csv('D08_SV메트릭.csv', ['SV_NAME', 'TABLE', 'METRIC', 'DATA_TYPE', 'EXPRESSION', 'SYNONYMS', 'NON_ADDITIVE_DIMENSIONS', 'COMMENT'],
              [[m['SEMANTIC_VIEW_NAME'], m['TABLE_NAME'], m['METRIC_NAME'], m['DATA_TYPE'], cut(m['EXPRESSION'], 400),
                cut(m['SYNONYMS'], 200), cut(m['NON_ADDITIVE_DIMENSIONS'], 100), cut(m['COMMENT'], 300)] for m in mets], foot)

    # ---- D09 ----
    write_csv('D09_Agent도구.csv', ['AGENT', 'TOOL_NAME', 'TOOL_TYPE', 'TARGET', 'DESCRIPTION(앞 200자)'], tools, foot)

    # ---- D01 (선택 · B2 캡처가 있을 때만) ----
    b2 = os.path.join(env, 'B2_Q01_1.csv')
    n_d01 = 0
    if os.path.exists(b2):
        rows = []
        for r in rd(b2):
            state = r['COVERAGE'] or ('(산식 미등록)' if r['HAS_FORMULA'] != 'Y' else '(대응 상태 미기재)')
            rows.append([r['DOC_ID'], r['DOC_ID'].split('-')[0], r['TITLE'], cut(r['SYNONYMS'], 150), cut(r['DEFINITION'], 300),
                         cut(r['FORMULA'], 300), r['AGENT_SCOPE'], r['TOOL_NAME'], cut(r['SV_FIELD'], 200), state,
                         r['IS_ADDITIVE'], r['SOURCE_FILE']])
        write_csv('D01_지표매핑.csv', ['DOC_ID', 'SOURCE_GROUP', 'METRIC', 'SYNONYMS', 'DEFINITION', 'FORMULA', 'AGENT_SCOPE',
                                     'TOOL_NAME', 'SV_FIELD', 'COVERAGE', 'IS_ADDITIVE', 'SOURCE_FILE'], rows,
                  foot.replace('B1_*', 'B1_*·B2_*'))
        n_d01 = len(rows)

    print('🟢 B · account=%s · 테이블 %d · 컬럼 %d · FK %d · 계보 간선 %d(미해결 %d · 물리 없는 dbt 모델 %d) · SV %d · 메트릭 %d · Agent 도구 %d · D01 지표 %d'
          % (acct, len(trows), len(crows), len(fk_rows), len(lin), len(broken), len(nophys), len(sv_rows), len(mets), len(tools), n_d01))
    if nophys:
        print('   물리 없는 dbt 모델:', ', '.join(nophys[:20]))
    if broken:
        print('   미해결 간선 예:', broken[:5])
    return 0


if __name__ == '__main__':
    sys.exit(main())
